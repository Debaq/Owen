import os
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

DIR = '/home/nick/Escritorio/Proyectos/Owen/screenshots'

doc = Document()

# ============================================================
# CONFIGURACIÓN DE ESTILOS
# ============================================================
sections = doc.sections
for section in sections:
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

style_normal = doc.styles['Normal']
style_normal.font.name = 'Calibri'
style_normal.font.size = Pt(11)
style_normal.font.color.rgb = RGBColor(0x2d, 0x2d, 0x2d)
style_normal.paragraph_format.space_after = Pt(6)
style_normal.paragraph_format.line_spacing = 1.15

AZUL_OSCURO = RGBColor(0x1a, 0x1a, 0x2e)
AZUL_MEDIO = RGBColor(0x2d, 0x3a, 0x8c)
AZUL_CLARO = RGBColor(0x3b, 0x82, 0xf6)
GRIS = RGBColor(0x6b, 0x72, 0x80)
VERDE = RGBColor(0x05, 0x96, 0x69)
ROJO = RGBColor(0xdc, 0x26, 0x26)
NARANJA = RGBColor(0xea, 0x58, 0x0c)

def heading(text, level=1):
    h = doc.add_heading(text, level=level)
    for run in h.runs:
        if level == 1:
            run.font.color.rgb = AZUL_OSCURO
            run.font.size = Pt(20)
        elif level == 2:
            run.font.color.rgb = AZUL_MEDIO
            run.font.size = Pt(15)
        elif level == 3:
            run.font.color.rgb = AZUL_CLARO
            run.font.size = Pt(12)
    return h

def para(text, bold=False, italic=False, size=11, color=None, align=None, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.size = Pt(size)
    run.font.name = 'Calibri'
    if bold:
        run.bold = True
    if italic:
        run.italic = True
    if color:
        run.font.color.rgb = color
    if align == 'center':
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == 'right':
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.paragraph_format.space_after = Pt(space_after)
    return p

def rich_para(parts, space_after=6):
    """parts: list of (text, {bold, italic, color, size})"""
    p = doc.add_paragraph()
    for text, fmt in parts:
        run = p.add_run(text)
        run.font.name = 'Calibri'
        run.font.size = Pt(fmt.get('size', 11))
        if fmt.get('bold'):
            run.bold = True
        if fmt.get('italic'):
            run.italic = True
        if fmt.get('color'):
            run.font.color.rgb = fmt['color']
    p.paragraph_format.space_after = Pt(space_after)
    return p

def bullet(text, level=0):
    style = 'List Bullet' if level == 0 else f'List Bullet {level + 1}'
    p = doc.add_paragraph(style=style)
    # Parse bold
    import re
    parts = re.split(r'(\*\*.*?\*\*)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = p.add_run(part[2:-2])
            run.bold = True
        else:
            p.add_run(part)
    return p

def add_image(filename, caption=None, width=Inches(5.5)):
    img_path = os.path.join(DIR, filename)
    if os.path.exists(img_path):
        doc.add_picture(img_path, width=width)
        last_paragraph = doc.paragraphs[-1]
        last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if caption:
            p = para(caption, italic=True, size=9, color=GRIS, align='center', space_after=12)
    else:
        para(f'[Imagen no encontrada: {filename}]', italic=True, color=ROJO)

def add_styled_table(headers, rows, col_widths=None):
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    # Header row
    for j, h in enumerate(headers):
        cell = table.cell(0, j)
        cell.text = h
        for p in cell.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.bold = True
                run.font.size = Pt(10)
                run.font.color.rgb = RGBColor(0xff, 0xff, 0xff)
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="1a1a2e"/>')
        cell._tc.get_or_add_tcPr().append(shading)

    # Data rows
    for i, row in enumerate(rows):
        for j, val in enumerate(row):
            cell = table.cell(i + 1, j)
            cell.text = str(val)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(10)
            bg = 'f8f9fa' if i % 2 == 0 else 'ffffff'
            shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{bg}"/>')
            cell._tc.get_or_add_tcPr().append(shading)

    doc.add_paragraph()

def separator():
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run('─' * 70)
    run.font.color.rgb = RGBColor(0xd1, 0xd5, 0xdb)
    run.font.size = Pt(8)
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER


# ============================================================
# PORTADA
# ============================================================
for _ in range(6):
    doc.add_paragraph()

para('SISTEMA OWEN', bold=True, size=36, color=AZUL_OSCURO, align='center', space_after=4)
para('Sistema de Gestión de Horarios Institucional', size=16, color=AZUL_MEDIO, align='center', space_after=24)

separator()

para('Documento de Proyecto', bold=True, size=14, color=GRIS, align='center', space_after=8)
para('Sede Puerto Montt, Chile', size=12, color=GRIS, align='center', space_after=24)

for _ in range(4):
    doc.add_paragraph()

add_styled_table(
    ['', ''],
    [
        ['Fecha', '30 de marzo de 2026'],
        ['Versión', '1.0'],
        ['Estado', 'MVP en producción'],
        ['URL', 'https://tmeduca.org/owen/'],
    ]
)

doc.add_page_break()

# ============================================================
# ÍNDICE
# ============================================================
heading('Índice de Contenidos', 1)
toc_items = [
    ('1.', 'Contexto y Problemática'),
    ('2.', 'Objetivos del Proyecto'),
    ('3.', 'Usuarios y Roles'),
    ('4.', 'Descripción Funcional'),
    ('5.', 'Portal Público'),
    ('6.', 'Panel de Administración'),
    ('7.', 'Arquitectura Técnica'),
    ('8.', 'Estado Actual del Proyecto'),
    ('9.', 'Hoja de Ruta'),
]
for num, title in toc_items:
    rich_para([
        (num + ' ', {'bold': True, 'color': AZUL_MEDIO, 'size': 12}),
        (title, {'size': 12}),
    ])

doc.add_page_break()

# ============================================================
# 1. CONTEXTO Y PROBLEMÁTICA
# ============================================================
heading('1. Contexto y Problemática', 1)

para('El campus universitario de Puerto Montt enfrenta desafíos significativos en la gestión de sus horarios y espacios físicos:')

bullet('**Gestión manual:** El sistema actual (Darwin) es tan deficiente que la asignación de salas, horarios y docentes se realiza de forma completamente manual, mediante planillas y coordinación informal.')
bullet('**Campus disperso:** Los edificios están distribuidos en un área extensa, con distancias de hasta 40 minutos a pie entre ellos. Con solo 10 minutos entre bloques horarios, la ubicación de las clases es crítica.')
bullet('**Condiciones climáticas:** Puerto Montt es una de las ciudades más lluviosas de Chile, lo que agrava el problema de los traslados entre edificios.')
bullet('**Sin visibilidad pública:** Estudiantes y visitantes no tienen forma de consultar horarios ni ubicar salas de forma autónoma.')
bullet('**Sin trazabilidad:** No existe un sistema de solicitudes formales para el uso de salas, ni seguimiento de observaciones o problemas de infraestructura.')

doc.add_paragraph()
para('Sistema OWEN nace como respuesta integral a estos problemas, reemplazando los procesos manuales con una plataforma digital que considera las particularidades geográficas y climáticas del campus.', italic=True, color=AZUL_MEDIO)

doc.add_page_break()

# ============================================================
# 2. OBJETIVOS
# ============================================================
heading('2. Objetivos del Proyecto', 1)

heading('Objetivo General', 2)
para('Desarrollar una plataforma web que centralice la gestión de horarios, salas y recursos académicos del campus, optimizando la asignación de espacios y mejorando la experiencia de todos los actores involucrados.')

heading('Objetivos Específicos', 2)
bullet('**Digitalizar la gestión de horarios:** Reemplazar las planillas manuales con un sistema web que permita crear, modificar y visualizar horarios en tiempo real.')
bullet('**Optimizar la asignación de salas:** Considerar capacidad, equipamiento, ubicación geográfica y distancias entre edificios para minimizar traslados.')
bullet('**Habilitar autoservicio público:** Permitir a estudiantes y visitantes consultar horarios y ubicar salas desde cualquier dispositivo, sin necesidad de autenticación.')
bullet('**Automatizar solicitudes:** Implementar un flujo inteligente de solicitudes de sala con asistencia de IA para auto-aprobación cuando la confianza es alta.')
bullet('**Gestionar infraestructura vía QR:** Permitir reportes anónimos de problemas en salas y áreas comunes mediante códigos QR instalados en el campus.')
bullet('**Generar horarios automáticamente:** Integrar un solver de optimización (HiGHS) que proponga distribuciones óptimas respetando restricciones académicas y físicas.')

doc.add_page_break()

# ============================================================
# 3. USUARIOS Y ROLES
# ============================================================
heading('3. Usuarios y Roles', 1)

add_styled_table(
    ['Rol', 'Acceso', 'Responsabilidades'],
    [
        ['Gestor', 'Panel completo', 'CRUD de horarios, salas, docentes, asignaturas.\nAprueba/rechaza solicitudes.\nGestiona observaciones y bloqueos.'],
        ['Dirección de Carrera', 'Panel limitado', 'Consulta horarios de su carrera (solo lectura).\nCrea solicitudes de sala.\nLibera clases cuando no se dictarán.'],
        ['Público General', 'Portal público', 'Consulta horarios por sala o carrera.\nExplora el mapa del campus.\nReporta observaciones anónimas vía QR.'],
    ]
)

doc.add_page_break()

# ============================================================
# 4. DESCRIPCIÓN FUNCIONAL
# ============================================================
heading('4. Descripción Funcional', 1)

para('El sistema se organiza en módulos que cubren tres grandes áreas:')

heading('Gestión Física', 2)
bullet('**Edificios:** Registro de edificios del campus con ubicación georreferenciada, cantidad de pisos, fotos y descripción.')
bullet('**Salas:** Catálogo completo de espacios (aulas, laboratorios, auditorios, talleres, oficinas, bibliotecas) con capacidad, equipamiento, tipo de mobiliario y nivel de gestión.')
bullet('**Mapa Interactivo:** Visualización del campus sobre OpenStreetMap con marcadores por edificio, estado de ocupación en tiempo real y trazado de rutas peatonales.')

heading('Gestión Académica', 2)
bullet('**Carreras y Niveles:** Estructura jerárquica de carreras → niveles → asignaturas.')
bullet('**Docentes:** Registro con RUT, unidad académica, carreras asociadas, disponibilidad horaria y carga académica.')
bullet('**Asignaturas:** Materias con horas de teoría, práctica y trabajo autónomo, divididas en sesiones y secciones.')

heading('Gestión de Horarios', 2)
bullet('**Grilla Semanal:** Visualización y edición de horarios por sala, docente, nivel o asignatura.')
bullet('**Asistente (Wizard):** Creación guiada paso a paso: sala → asignatura → docente → bloque → recurrencia.')
bullet('**Detección de Conflictos:** Verificación automática de sala ocupada, docente con clase simultánea, nivel con clase simultánea, sala bloqueada o feriado.')
bullet('**Bloques Horarios:** Sistemas configurables de franjas horarias (vespertino, diurno, etc.) con generación masiva.')
bullet('**Versionado:** Sistema planificado tipo Git (branches, commits, diff, merge) para gestionar borradores y versiones de horarios.')

heading('Módulos de Soporte', 2)
bullet('**Solicitudes de Sala:** Flujo formal de peticiones con análisis de disponibilidad asistido por IA (Claude API) y auto-aprobación.')
bullet('**Observaciones vía QR:** Reporte anónimo de problemas en infraestructura con ciclo de vida completo (nuevo → revisión → en proceso → resuelto → cerrado).')
bullet('**Reportes:** Exportación de horarios en PDF y Excel, filtrados por sala, docente, nivel o asignatura.')
bullet('**Solver:** Aplicación de escritorio (Tauri + Rust) que se conecta a Owen vía API para generar horarios óptimos con el motor HiGHS.')

doc.add_page_break()

# ============================================================
# 5. PORTAL PÚBLICO
# ============================================================
heading('5. Portal Público', 1)

para('El portal público es accesible sin autenticación y está diseñado para que cualquier persona pueda consultar información del campus.')

heading('Página Principal', 2)
para('Presenta un buscador de salas con dos modos de consulta: por sala (texto libre + filtros por tipo) y por carrera/nivel (selectores jerárquicos). Incluye acceso a lugares de interés y un mapa interactivo del campus.')

add_image('01_home.png', 'Figura 1 — Portal público: buscador de salas y mapa del campus')

heading('Búsqueda de Salas', 2)
para('Los usuarios pueden buscar salas por nombre, código o edificio, y filtrar por tipo de espacio: Aula, Laboratorio, Auditorio, Taller, Sala de Reuniones, Oficina, Biblioteca o Medioteca.')

add_image('04_tab_por_carrera.png', 'Figura 2 — Búsqueda por carrera y nivel académico')

heading('Mapa del Campus', 2)
para('Mapa interactivo basado en Leaflet y OpenStreetMap que muestra los edificios del campus con indicadores de estado (disponible/ocupada). Los usuarios pueden explorar el campus, ubicar edificios y consultar salas desde el mapa.')

add_image('05_explorar_campus.png', 'Figura 3 — Explorador del campus con lugares de interés')

heading('Inicio de Sesión', 2)
para('Formulario de autenticación limpio y accesible desde el botón "Iniciar Sesión" del portal o la ruta directa /login.')

add_image('07_login_page.png', 'Figura 4 — Formulario de inicio de sesión', width=Inches(3.5))

doc.add_page_break()

# ============================================================
# 6. PANEL DE ADMINISTRACIÓN
# ============================================================
heading('6. Panel de Administración', 1)

para('El panel de administración está organizado en un sidebar lateral con tres categorías: Gestión Física, Gestión Académica y Sistema.')

heading('Dashboard', 2)
para('Vista consolidada con contadores de entidades principales (salas, edificios, docentes, carreras, horarios), ocupación del día en tiempo real, solicitudes recientes, últimos horarios creados y accesos rápidos a las funciones más utilizadas.')

add_image('01_dashboard.png', 'Figura 5 — Dashboard del panel de administración')

heading('Gestión de Horarios', 2)
para('Grilla semanal interactiva con filtros por sala, docente, nivel o asignatura. Cada celda permite agregar un horario directamente. La leyenda identifica asignatura, docente y sala con colores.')

add_image('02_horarios.png', 'Figura 6 — Grilla semanal de horarios')

heading('Asistente de Horarios', 2)
para('Wizard guiado para la creación de horarios en pasos secuenciales. El primer paso permite seleccionar la sala mostrando código, tipo, nombre y capacidad, con opción de crear una nueva sala sobre la marcha.')

add_image('04_asistente_horarios.png', 'Figura 7 — Asistente de creación de horarios (paso 1: selección de sala)')

heading('Edificios', 2)
para('Catálogo visual de edificios con tarjetas que muestran código, descripción, cantidad de pisos y aulas. El formulario de creación incluye ubicación georreferenciada seleccionable en el mapa.')

add_image('05_edificios.png', 'Figura 8 — Catálogo de edificios del campus')
add_image('17_edificio_crear_form.png', 'Figura 9 — Formulario de registro de edificio con mapa de ubicación')

heading('Gestión de Salas', 2)
para('Tarjetas por sala con información de tipo, capacidad, equipamiento (proyector, pizarra), nivel de gestión y acceso rápido a QR y vista pública. El formulario de creación es completo con todos los atributos configurables.')

add_image('06_salas.png', 'Figura 10 — Catálogo de salas con estado y equipamiento')
add_image('18_sala_crear_form.png', 'Figura 11 — Formulario de creación de sala')

heading('Mapa del Campus (Administración)', 2)
para('Vista administrativa del mapa con herramientas adicionales para gestionar marcadores, edificios y puntos de interés. Incluye leyenda de estado de salas y control de capas.')

add_image('07_mapa_admin.png', 'Figura 12 — Mapa del campus en modo administración')

doc.add_page_break()

heading('Gestión Académica', 2)

para('Las secciones de Carreras, Unidades y Docentes comparten un diseño master-detail: lista con búsqueda en el panel izquierdo y detalle completo en el panel derecho.')

add_image('08_carreras.png', 'Figura 13 — Gestión de carreras (vista master-detail)')
add_image('10_docentes.png', 'Figura 14 — Gestión de docentes con perfil y carga académica')

heading('Bloques Horarios', 2)
para('Configuración de sistemas de bloques (vespertino, diurno, etc.) con grilla visual de lunes a viernes. Cada bloque muestra su rango horario exacto. Incluye generador masivo para crear todos los bloques de un sistema de una vez.')

add_image('11_bloques_horarios.png', 'Figura 15 — Configuración de bloques horarios del sistema vespertino')

heading('Solicitudes y Observaciones', 2)
para('Módulo de solicitudes de sala con flujo de aprobación/rechazo. Las observaciones permiten gestionar reportes de infraestructura enviados vía QR desde el campus.')

add_image('12_solicitudes.png', 'Figura 16 — Gestión de solicitudes de sala')

heading('Reportes', 2)
para('Generación de reportes exportables a PDF y Excel, con filtros por sala, docente, nivel o asignatura y vista previa antes de la exportación.')

add_image('14_reportes.png', 'Figura 17 — Módulo de reportes con exportación PDF/Excel')

heading('Configuración del Sistema', 2)
para('Panel de configuración con pestañas para identidad del sitio, notificaciones, área pública y generación de QR. Incluye integración con OpenRouteService para trazado de rutas peatonales y control del módulo Solver.')

add_image('15_configuracion.png', 'Figura 18 — Configuración general del sistema')

doc.add_page_break()

# ============================================================
# 7. ARQUITECTURA TÉCNICA
# ============================================================
heading('7. Arquitectura Técnica', 1)

heading('Stack Tecnológico', 2)

add_styled_table(
    ['Capa', 'Tecnología', 'Justificación'],
    [
        ['Frontend', 'React 19 + TypeScript + Vite', 'Rendimiento, tipado estático, recarga rápida en desarrollo'],
        ['UI', 'Tailwind CSS + Shadcn/ui', 'Diseño consistente, componentes accesibles y personalizables'],
        ['Mapas', 'React-Leaflet + OpenStreetMap', 'Mapas gratuitos, sin dependencia de APIs propietarias'],
        ['Rutas', 'OpenRouteService', 'Trazado de rutas peatonales dentro del campus'],
        ['Backend', 'PHP 7.4 (API REST JSON)', 'Compatible con infraestructura existente del servidor'],
        ['Base de datos', 'SQLite3', 'Sin dependencias externas, portable, suficiente para la escala'],
        ['Autenticación', 'Sesiones PHP + cookies seguras', 'Estándar, sin complejidad adicional'],
        ['IA', 'Claude API', 'Análisis inteligente de solicitudes de sala'],
        ['Solver', 'Tauri + Rust + HiGHS', 'App de escritorio para optimización de horarios'],
        ['Internacionalización', 'i18next (ES/EN)', 'Soporte bilingüe'],
    ]
)

heading('Estructura del Proyecto', 2)

add_styled_table(
    ['Directorio', 'Contenido'],
    [
        ['src/features/', 'Módulos funcionales (auth, buildings, rooms, schedules, map, academic, settings, etc.)'],
        ['src/shared/', 'Componentes UI reutilizables, tipos TypeScript, hooks y utilidades'],
        ['backend/api/', 'Endpoints REST PHP (auth, salas, edificios, horarios, docentes, etc.)'],
        ['backend/db/', 'Base de datos SQLite y esquema SQL'],
    ]
)

doc.add_page_break()

# ============================================================
# 8. ESTADO ACTUAL
# ============================================================
heading('8. Estado Actual del Proyecto', 1)

para('El sistema se encuentra en estado de MVP funcional desplegado en producción. A continuación se detalla el estado de cada módulo:')

heading('Módulos Operativos', 2)

add_styled_table(
    ['Módulo', 'Estado', 'Observaciones'],
    [
        ['Autenticación', '✓ Completo', 'Login/logout con sesiones, roles gestor y dirección'],
        ['Dashboard', '✓ Completo', 'Contadores, widgets de ocupación, accesos rápidos'],
        ['Edificios', '✓ Completo', 'CRUD completo con fotos y geolocalización'],
        ['Salas', '✓ Completo', 'CRUD con equipamiento, QR, gestión y vista pública'],
        ['Mapa del Campus', '✓ Completo', 'Interactivo con marcadores, capas y herramientas admin'],
        ['Carreras', '✓ Completo', 'Master-detail con niveles y asignaturas'],
        ['Docentes', '✓ Completo', 'Master-detail con disponibilidad y carga académica'],
        ['Bloques Horarios', '✓ Completo', 'Múltiples sistemas, generador masivo, temporadas'],
        ['Grilla de Horarios', '✓ Completo', 'Vista semanal por sala, docente, nivel o asignatura'],
        ['Asistente de Horarios', '✓ Completo', 'Wizard paso a paso para creación guiada'],
        ['Reportes', '✓ Completo', 'Exportación PDF/Excel con filtros y vista previa'],
        ['Configuración', '✓ Completo', 'Identidad, mapa, módulos, migraciones'],
        ['Portal Público', '✓ Completo', 'Búsqueda de salas, mapa, i18n'],
    ]
)

heading('Módulos Pendientes', 2)

add_styled_table(
    ['Módulo', 'Estado', 'Descripción'],
    [
        ['Observaciones QR', '○ Pendiente', 'Formulario público y panel de gestión de tickets'],
        ['Solicitudes con IA', '◐ Parcial', 'UI lista, falta integración Claude API para auto-aprobación'],
        ['Versionado de Horarios', '○ Pendiente', 'Modelo tipo Git para borradores y versiones'],
        ['Detección de Conflictos', '○ Pendiente', 'Validación de sala, docente y nivel simultáneo'],
        ['Rol Dirección', '○ Pendiente', 'Vistas limitadas y liberación de clases'],
        ['Calendario/Feriados', '○ Pendiente', 'Visualización y gestión de días no lectivos'],
        ['Solver (HiGHS)', '◐ Parcial', 'App Tauri separada, integración API pendiente'],
        ['Links Públicos', '○ Pendiente', 'URLs compartibles para horarios específicos'],
    ]
)

doc.add_page_break()

# ============================================================
# 9. HOJA DE RUTA
# ============================================================
heading('9. Hoja de Ruta', 1)

heading('Fase 1 — Completar el núcleo', 2)
para('Prioridad: Alta', bold=True, color=ROJO)
bullet('Implementar detección de conflictos al crear/editar horarios')
bullet('Completar el módulo de observaciones vía QR (formulario público + gestión admin)')
bullet('Activar el flujo completo de solicitudes de sala con notificaciones')
bullet('Implementar el rol Dirección con vistas limitadas y liberación de clases')
bullet('Poblar la base de datos con datos reales del campus')

heading('Fase 2 — Inteligencia y optimización', 2)
para('Prioridad: Media', bold=True, color=NARANJA)
bullet('Integrar Claude API para análisis y auto-aprobación de solicitudes')
bullet('Conectar el Solver (Tauri/HiGHS) con Owen vía API para generación automática de horarios')
bullet('Implementar versionado de horarios (branches, commits, diff, merge)')
bullet('Agregar calendario con feriados y eventos institucionales')

heading('Fase 3 — Experiencia y alcance', 2)
para('Prioridad: Normal', bold=True, color=VERDE)
bullet('Generar links públicos compartibles para horarios específicos')
bullet('Implementar notificaciones (email/push) ante cambios de horario')
bullet('Optimizar la experiencia móvil del portal público')
bullet('Documentar API para integraciones con otros sistemas institucionales')

doc.add_paragraph()
separator()
doc.add_paragraph()

para('Sistema OWEN — Sede Puerto Montt, Chile', italic=True, size=10, color=GRIS, align='center')
para('Documento generado el 30 de marzo de 2026', italic=True, size=9, color=GRIS, align='center')

# ============================================================
# GUARDAR
# ============================================================
out = os.path.join(DIR, 'DOCUMENTO_PROYECTO_OWEN.docx')
doc.save(out)
print(f'✅ Documento guardado: {out}')
print(f'   Tamaño: {os.path.getsize(out) / 1024:.0f} KB')
