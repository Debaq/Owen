import re
import sys
from docx import Document
from docx.shared import Pt, Inches, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

doc = Document()

# Estilos
style = doc.styles['Normal']
font = style.font
font.name = 'Calibri'
font.size = Pt(11)

style_h1 = doc.styles['Heading 1']
style_h1.font.size = Pt(22)
style_h1.font.color.rgb = RGBColor(0x1a, 0x1a, 0x2e)

style_h2 = doc.styles['Heading 2']
style_h2.font.size = Pt(16)
style_h2.font.color.rgb = RGBColor(0x2d, 0x3a, 0x8c)

style_h3 = doc.styles['Heading 3']
style_h3.font.size = Pt(13)
style_h3.font.color.rgb = RGBColor(0x3b, 0x82, 0xf6)

with open('/home/nick/Escritorio/Proyectos/Owen/screenshots/DOCUMENTO_PROYECTO_OWEN.md', 'r') as f:
    lines = f.readlines()

def add_table(doc, rows):
    if not rows:
        return
    headers = rows[0]
    table = doc.add_table(rows=len(rows), cols=len(headers))
    table.style = 'Light Grid Accent 1'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(rows):
        for j, cell in enumerate(row):
            table.cell(i, j).text = cell.strip()
            for p in table.cell(i, j).paragraphs:
                p.style.font.size = Pt(10)
    doc.add_paragraph()

def parse_inline(paragraph, text):
    """Parsea bold, italic, code inline"""
    parts = re.split(r'(\*\*.*?\*\*|\*.*?\*|`[^`]+`)', text)
    for part in parts:
        if part.startswith('**') and part.endswith('**'):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
        elif part.startswith('*') and part.endswith('*'):
            run = paragraph.add_run(part[1:-1])
            run.italic = True
        elif part.startswith('`') and part.endswith('`'):
            run = paragraph.add_run(part[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(9)
            run.font.color.rgb = RGBColor(0x99, 0x33, 0x33)
        else:
            paragraph.add_run(part)

i = 0
in_code_block = False
code_lines = []
in_table = False
table_rows = []

while i < len(lines):
    line = lines[i].rstrip('\n')

    # Code blocks
    if line.startswith('```'):
        if in_code_block:
            # End code block
            p = doc.add_paragraph()
            p.style = doc.styles['Normal']
            for cl in code_lines:
                run = p.add_run(cl + '\n')
                run.font.name = 'Consolas'
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
            code_lines = []
            in_code_block = False
        else:
            in_code_block = True
        i += 1
        continue

    if in_code_block:
        code_lines.append(line)
        i += 1
        continue

    # Tables
    if '|' in line and line.strip().startswith('|'):
        cells = [c.strip() for c in line.strip().strip('|').split('|')]
        # Skip separator rows
        if all(re.match(r'^[-:]+$', c) for c in cells):
            i += 1
            continue
        table_rows.append(cells)
        # Check if next line is not a table row
        if i + 1 >= len(lines) or '|' not in lines[i + 1] or not lines[i + 1].strip().startswith('|'):
            add_table(doc, table_rows)
            table_rows = []
        i += 1
        continue

    # Horizontal rule
    if line.strip() == '---':
        doc.add_paragraph('─' * 60)
        i += 1
        continue

    # Headers
    if line.startswith('# '):
        doc.add_heading(line[2:], level=1)
        i += 1
        continue
    if line.startswith('## '):
        doc.add_heading(line[3:], level=2)
        i += 1
        continue
    if line.startswith('### '):
        doc.add_heading(line[4:], level=3)
        i += 1
        continue

    # Empty line
    if line.strip() == '':
        i += 1
        continue

    # List items
    if line.startswith('- '):
        p = doc.add_paragraph(style='List Bullet')
        parse_inline(p, line[2:])
        i += 1
        continue
    if re.match(r'^  - ', line):
        p = doc.add_paragraph(style='List Bullet 2')
        parse_inline(p, line[4:])
        i += 1
        continue
    if re.match(r'^    - ', line):
        p = doc.add_paragraph(style='List Bullet 3')
        parse_inline(p, line[6:])
        i += 1
        continue

    # Numbered list
    m = re.match(r'^(\d+)\. ', line)
    if m:
        p = doc.add_paragraph(style='List Number')
        parse_inline(p, line[len(m.group(0)):])
        i += 1
        continue

    # Normal paragraph
    p = doc.add_paragraph()
    parse_inline(p, line)
    i += 1

out = '/home/nick/Escritorio/Proyectos/Owen/screenshots/DOCUMENTO_PROYECTO_OWEN.docx'
doc.save(out)
print(f'Guardado en {out}')
